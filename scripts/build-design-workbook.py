"""Generate an Indonesian, fillable PDF brief and an editable Word companion.

Requirements: reportlab, python-docx, PyMuPDF (only to export guide illustrations).
Run from any directory: python scripts/build-design-workbook.py
"""
from pathlib import Path
from collections import Counter
from xml.sax.saxutils import escape
import math
import json

from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fitz

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/design'
ASSETS = OUT / 'brief-assets'
STEM = 'Elemonts-Formulir-Kebutuhan-Visual-dan-Battle-v1'
PDF = OUT / (STEM + '.pdf')
DOCX = OUT / (STEM + '.docx')
SECTIONS = []
for line in (ASSETS / 'questions.txt').read_text().splitlines():
    if not line or line.startswith('#'):
        continue
    cells = [x.strip() for x in line.split('|')]
    if cells[0] == 'SECTION':
        _, num, title, owner, priority = cells
        SECTIONS.append(dict(num=num, title=title, owner=owner, priority=priority, questions=[]))
    else:
        assert len(cells) == 6, cells
        SECTIONS[-1]['questions'].append(dict(zip(['id', 'question', 'meaning', 'options', 'example', 'evidence'], cells)))
QUESTIONS = [q for s in SECTIONS for q in s['questions']]
assert len(QUESTIONS) == 72

W, H = A4
M = 40
CW = W - 2*M
INK = HexColor('#17333D')
MUTED = HexColor('#526B72')
TEAL = HexColor('#087F80')
PALE = HexColor('#EFF8F8')
GOLD = HexColor('#B36A19')
EXAMPLE = HexColor('#FFF5E7')
LINE = HexColor('#CCDADD')
PAPER = HexColor('#FCFDFC')
FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
BOLD = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
pdfmetrics.registerFont(TTFont('Brief', FONT))
pdfmetrics.registerFont(TTFont('BriefBold', BOLD))
pdfmetrics.registerFontFamily('Brief', normal='Brief', bold='BriefBold')

def paragraph(text, size=10, leading=None, color=INK, bold=False):
    return Paragraph(text, ParagraphStyle('p', fontName='BriefBold' if bold else 'Brief',
        fontSize=size, leading=leading or size*1.35, textColor=color, spaceAfter=0))

def para(c, text, x, top, width, size=10, color=INK, bold=False):
    p = paragraph(text, size, color=color, bold=bold)
    _, h = p.wrap(width, 1000)
    p.drawOn(c, x, H-top-h)
    return top+h

def box(c, x, top, width, height, fill=PALE, stroke=LINE, radius=8):
    c.setFillColor(fill)
    c.setStrokeColor(stroke)
    c.roundRect(x,H-top-height,width,height,radius,fill=1,stroke=1)

def field(c, name, x, top, width, height, multiline=False, value=''):
    c.acroForm.textfield(name=name, tooltip=name.replace('_',' '),
        x=x, y=H-top-height, width=width, height=height,
        value=value, fontName='Helvetica', fontSize=10,
        borderWidth=0.7, borderColor=TEAL, fillColor=white, textColor=INK,
        forceBorder=True, fieldFlags='multiline' if multiline else '', maxlen=6000)

page = 0
TOTAL = 5 + sum(math.ceil(len(s['questions'])/2) for s in SECTIONS) + 3
c = canvas.Canvas(str(PDF), pagesize=A4)
c.setTitle('Elemonts | Formulir Kebutuhan Visual dan Battle')
c.setAuthor('Tim Elemonts / Codex')
c.setSubject('72 pertanyaan, petunjuk, contoh, referensi, dan formulir isian')

def start(title, subtitle='', bookmark=None):
    global page
    if page:
        c.showPage()
    page += 1
    c.setFillColor(PAPER); c.rect(0,0,W,H,fill=1,stroke=0)
    c.setFillColor(TEAL); c.rect(0,H-8,W,8,fill=1,stroke=0)
    para(c, 'ELEMONTS  /  BRIEF PRODUKSI', M, 23, CW, 8, TEAL, True)
    bottom = para(c, title, M, 43, CW, 19, bold=True)
    if subtitle:
        para(c, subtitle, M, bottom+7, CW, 9.3, MUTED)
    c.setStrokeColor(LINE); c.line(M,40,W-M,40)
    para(c, 'v1.0 • 25 Oktober 2026 • Contoh bukan keputusan final', M, H-31, CW-70, 7.2, MUTED)
    para(c, f'{page:02d} / {TOTAL:02d}', W-M-50,H-32,50,8,TEAL,True)
    if bookmark:
        c.bookmarkPage(bookmark)
        c.addOutlineEntry(title,bookmark,level=0)

start('Formulir kebutuhan\nvisual dan battle'.replace('\n','<br/>'),
      'Panduan kerja untuk pemilik game, tim desain, artist, animator, dan QA.', 'cover')
box(c,M,144,CW,111,INK,INK)
para(c,'Dari referensi menjadi keputusan yang bisa dikerjakan.',M+18,159,CW-36,17,white,True)
para(c,'72 pertanyaan • 12 bagian • Kolom PDF dapat diketik<br/>Contoh kasus • Panduan gambar/video • Versi Word untuk menempel gambar',M+18,213,CW-36,10,white)
y=275
for title, body in [
 ('Apa yang sedang ditingkatkan?', 'Kualitas bentuk monster, animasi tubuh, kamera, efek serangan, arena, suara, dan kenyamanan battle Android Elemonts.'),
 ('Apa yang tetap menjadi acuan?', 'Dokumen konsep dan aturan yang sudah disepakati, roster monster, serta Excel bentuk monster. File bentuk digunakan untuk tampilan, bukan mengganti aturan elemen.'),
 ('Apa yang belum menjadi keputusan?', 'Pilihan dan rekomendasi di formulir ini belum disetujui otomatis. Battle dua monster (usulan: Pyroo dan Croobit) adalah cara yang disarankan untuk menetapkan kualitas.'),
 ('Apakah semua harus dijawab sekarang?', 'Tidak. Jawab bagian prioritas AWAL dahulu. Untuk rincian visual yang belum diputuskan, tulis USULKAN atau BELUM TAHU; jangan biarkan kosong tanpa keterangan.')]:
    y=para(c,title,M,y,CW,11,bold=True)+5
    y=para(c,body,M,y,CW,10)+15
para(c,'IDENTITAS PENGISI — KETIK DI KOLOM DI BAWAH',M,650,CW,9,TEAL,True)
for idx,(label,key,x,width) in enumerate([
 ('Nama / peran','nama',M,250),('Tim / proyek','tim',M+265,CW-265)]):
    para(c,label,x,672,width,8.5,MUTED); field(c,key,x,688,width,24)
para(c,'Tanggal / versi jawaban',M,725,235,8.5,MUTED);field(c,'versi_jawaban',M,740,250,24)
para(c,'Gunakan halaman terakhir untuk ringkasan pengiriman.',M+265,735,CW-265,8.5,MUTED)

start('Cara membaca dan mengisi', 'Bullet pada daftar lama tidak semuanya jawaban. Di sini setiap fungsi diberi label.', 'guide')
labels=[
 ('MAKSUD','Penjelasan mengapa pertanyaan ditanyakan. Tidak perlu dijawab terpisah.'),
 ('CARA MENJAWAB / PILIHAN','Pilih satu, pilih beberapa, urutkan, atau jawab bebas sesuai instruksinya.'),
 ('CONTOH KASUS + JAWABAN','Contoh agar mudah dipahami; bukan jawaban kamu dan bukan keputusan proyek.'),
 ('JAWABAN KAMU / TIM','Satu-satunya kotak utama untuk menulis keputusan. Pilihan boleh memakai huruf + alasan.'),
 ('LAMPIRAN','Petunjuk perlu/tidaknya gambar, video, file, atau sumber. Gambar tidak wajib pada semua pertanyaan.'),
 ('REF / LINK / TIMESTAMP','Isi kode REF-01, nama file/link, dan waktu video. Gambar ditempel di lembar referensi, bukan di kolom teks ini.')]
y=108
for label,body in labels:
    box(c,M,y,CW,51,EXAMPLE if label.startswith('CONTOH') else PALE)
    para(c,label,M+11,y+7,CW-22,8.4,TEAL,True)
    para(c,body,M+11,y+21,CW-22,8.6)
    y+=58
y=para(c,'Cara memakai file',M,y+3,CW,11,bold=True)+8
y=para(c,'PDF: unduh, buka pada pembaca PDF yang mendukung formulir, klik kotak, ketik, lalu Simpan sebagai. Coba tutup dan buka kembali untuk memastikan isian tersimpan. Jika preview hanya menampilkan halaman, buka file di aplikasi PDF atau gunakan Word.',M,y,CW,9.2)+6
y=para(c,'Word: ganti teks [TULIS DI SINI] dan tempel gambar pada kotak [TEMPEL GAMBAR]. PDF tidak memiliki tombol unggah gambar; gunakan anotasi gambar jika aplikasimu mendukung, atau kirim gambar terpisah dengan kode REF.',M,y,CW,9.2)+12
y=para(c,'Peta bagian / klik judul untuk menuju bagian',M,y,CW,10,bold=True)+9
first=6
for i,s in enumerate(SECTIONS):
    col=i//6; row=i%6
    xx=M+col*(CW/2+4); yy=y+row*19
    text=f"{s['num']}. {s['title']}"
    # Keep long TOC items readable by using two columns with short labels.
    short=['Target dan arah visual','Batas konsep','Desain monster','Aset dan pengadaan','Gerakan monster','Skill dan efek','Kamera dan arena','UI dan kenyamanan','Musik dan suara','Perangkat Android','Cakupan dan waktu','Kriteria penilaian'][i]
    para(c,f"{s['num']}. {short}  ·  hlm. {first}",xx,yy,CW/2-8,8.5,TEAL)
    c.linkRect('',f"section-{s['num']}",(xx,H-yy-15,xx+CW/2-8,H-yy),relative=0,thickness=0)
    first+=math.ceil(len(s['questions'])/2)

start('Contoh pengisian yang benar', 'Halaman ini hanya demonstrasi. Jawaban sungguhan ditulis pada bagian bernomor.', 'example')
box(c,M,112,CW,214,EXAMPLE)
y=para(c,'CONTOH — BUKAN KEPUTUSAN ELEMONTS',M+14,127,CW-28,9,GOLD,True)+12
y=para(c,'7.2  Seberapa aktif kamera boleh bergerak?',M+14,y,CW-28,12,bold=True)+10
y=para(c,'Pilihan: A hampir diam; B bergerak ringan; C banyak pergantian sudut.',M+14,y,CW-28,10)+10
y=para(c,'Contoh jawaban: B untuk serangan biasa; close-up hanya saat skill spesial. Kamera kembali ke sudut utama setelah benturan.',M+14,y,CW-28,10)+10
para(c,'Contoh referensi: REF-01 | video Lumi Master yang dikirim | 00:18-00:20 | suka sorotan monster saat mengeluarkan skill; tidak meminta menyalin karakter.',M+14,y,CW-28,9.5)
y=350
for title,body in [
 ('Jika memilih satu','Tulis: "1.2: A — garis tepi tipis, warna cerah." Tidak perlu menyalin semua pilihan.'),
 ('Jika memilih beberapa atau mengurutkan','Tulis: "1.1: B > D > C — gerakan tubuh dahulu, lalu efek, lalu kamera."'),
 ('Jika belum punya keputusan','Tulis: "USULKAN; batasnya harus lancar di HP target." Developer menyiapkan opsi untuk ditinjau, bukan menganggap kolom kosong sebagai persetujuan.'),
 ('Jika tidak punya gambar','Tulis: "Belum ada gambar; pakai deskripsi Excel." Kekurangan gambar tidak otomatis menghentikan seluruh pekerjaan.'),
 ('Jika jawaban terlalu panjang','Tulis ringkasan di kolom, lalu "lanjutan: jawaban-3.4.docx". Kolom PDF dapat menggulir; teks yang melampaui kotak mungkin tidak seluruhnya tercetak.'),
 ('Jika tim berbeda pendapat','Tulis opsi masing-masing dan beri status DRAF. PIC keputusan akhir merangkum pilihan sebelum ditandai FINAL.')]:
    y=para(c,title,M,y,CW,10.5,bold=True)+4
    y=para(c,body,M,y,CW,9.5)+13

start('Visual: cara menjelaskan referensi', 'Cuplikan dari video Lumi Master yang kamu kirim. Bahan diskusi, bukan aset game Elemonts.', 'visual')
for i,(fn,label) in enumerate([('reference-18s.jpg','Sekitar 00:18 — arena dan posisi'),('reference-19-5s.jpg','Sekitar 00:19,5 — sorotan aksi')]):
    x=M+i*(CW/2+8)
    c.drawImage(str(ASSETS/fn),x,H-112-285,width=160.3125,height=285)
    para(c,label,x,408,CW/2-10,9,TEAL,True)
    desc=('Tandai komposisi monster, ruang arena, atau UI yang disukai. Satu gambar belum menjelaskan kecepatan gerak.' if i==0 else 'Sebutkan apakah yang disukai pose, kamera mendekat, atau efek api. Untuk gerakan, sertakan rentang waktu video.')
    para(c,desc,x,439,CW/2-14,9.3)
box(c,M,510,CW,91,EXAMPLE)
para(c,'CONTOH CATATAN REFERENSI',M+12,522,CW-24,9,GOLD,True)
para(c,'"REF-01, untuk 7.2 dan 6.5. Suka: kamera menyorot monster saat bersiap memakai skill. Jangan diikuti: model karakter dan UI persisnya. Terapkan: satu close-up singkat, kemudian kembali ke arena."',M+12,541,CW-24,10)
para(c,'Urutan aksi yang bisa dibicarakan bersama',M,622,CW,11,bold=True)
stages=[('01','BERSIAP','Pose tubuh'),('02','SERANG','Gerak / efek'),('03','KONTAK','Reaksi + HP'),('04','PULIH','Kembali siap')]
bw=(CW-24)/4
for i,(num,title,desc) in enumerate(stages):
    xx=M+i*(bw+8)
    box(c,xx,648,bw,69)
    para(c,num,xx+9,656,bw-18,12,TEAL,True)
    para(c,title,xx+9,678,bw-18,8.6,bold=True)
    para(c,desc,xx+9,695,bw-18,8,MUTED)
para(c,'Diagram ini menjelaskan tahap animasi; bukan durasi final, rumus damage, atau perubahan aturan battle.',M,735,CW,9,MUTED)

start('Visual: dari gambar ke aset bergerak', 'Ilustrasi skematis — bentuk di bawah hanya contoh, bukan desain monster yang disetujui.', 'assets')
def creature(cx,cy,rig=False):
    # Canvas coordinates; joint sketch distinguishes a drawing from an animation rig.
    c.setFillColor(HexColor('#E7F3EF'));c.setStrokeColor(TEAL);c.setLineWidth(2)
    c.ellipse(cx-31,cy-20,cx+31,cy+37,fill=1,stroke=1)
    c.circle(cx,cy+56,25,fill=1,stroke=1)
    c.line(cx-18,cy-16,cx-25,cy-43);c.line(cx+18,cy-16,cx+25,cy-43)
    c.line(cx-28,cy+20,cx-44,cy+3);c.line(cx+28,cy+20,cx+44,cy+3)
    if rig:
        c.setStrokeColor(GOLD);c.setFillColor(GOLD)
        for a,b,d,e in [(0,56,0,15),(0,15,-25,-43),(0,15,25,-43),(0,30,-44,3),(0,30,44,3)]:
            c.line(cx+a,cy+b,cx+d,cy+e)
        for a,b in [(0,56),(0,15),(0,30),(-25,-43),(25,-43),(-44,3),(44,3)]:
            c.circle(cx+a,cy+b,3,fill=1,stroke=0)
for i,(title,desc,rig) in enumerate([
 ('1. BENTUK','Gambar dan warna<br/>menentukan penampilan.',False),
 ('2. MODEL + RIG','Kerangka memungkinkan<br/>anggota tubuh bergerak.',True),
 ('3. ANIMASI','Pose dirangkai menjadi<br/>aksi dan reaksi.',True)]):
    xx=M+i*(CW/3+3)
    box(c,xx,113,CW/3-7,213)
    para(c,title,xx+10,126,CW/3-27,9.5,TEAL,True)
    creature(xx+(CW/3-7)/2,H-222,rig)
    para(c,desc,xx+10,286,CW/3-27,8.7)
y=345
for title,body in [
 ('Contoh kasus','Sebuah gambar monster tampak bagus tetapi tidak memiliki bagian tubuh yang bisa digerakkan. Untuk animasi mengayun tangan, diperlukan model, susunan sendi, dan gerakan yang sesuai; menambah partikel saja belum menghasilkan ayunan.'),
 ('Yang dapat diberikan pemilik game','Ciri wajib, proporsi, warna, ekspresi, gambar depan/samping/belakang jika ada, serta contoh gerakan yang disukai. Kamu tidak harus menentukan jumlah poligon atau nama tulang.'),
 ('Yang disiapkan tim aset/developer','Model dan tekstur, rig, animasi diam/serang/terkena/KO, titik keluarnya efek, kamera, suara, integrasi, dan pengukuran performa Android.'),
 ('Istilah singkat','Rig = kerangka penggerak. Tekstur = gambar warna/detail permukaan. VFX = efek visual. SFX = efek suara. UI = tombol dan informasi layar. FPS = jumlah gambar per detik. PIC = penanggung jawab. KO = monster tidak dapat melanjutkan bertarung.')]:
    y=para(c,title,M,y,CW,10.5,bold=True)+5
    y=para(c,body,M,y,CW,9.5)+16
box(c,M,704,CW,65,EXAMPLE)
para(c,'Belum memiliki aset? Itu jawaban yang sah.',M+12,715,CW-24,10,GOLD,True)
para(c,'Tulis BELUM ADA pada 4.1 dan pilih cara pengadaan pada 4.3. Contoh battle bisa dimulai dengan aset sementara; kualitas final ditinjau setelah aset tersedia.',M+12,735,CW-24,9)

layout_log=[]
def question_card(q, top, height=314):
    box(c,M,top,CW,height,white)
    x=M+13; width=CW-26
    y=para(c,f"{q['id']}  {escape(q['question'])}",x,top+10,width,11,bold=True)+7
    for label,text,color in [('MAKSUD',q['meaning'],INK),('CARA MENJAWAB / PILIHAN',q['options'],INK),('CONTOH KASUS + JAWABAN',q['example'],GOLD),('LAMPIRAN',q['evidence'],MUTED)]:
        y=para(c,f'<b>{label}:</b> {escape(text)}',x,y,width,9.2,color)+5
    # Reserve a predictable writing area at the bottom of every card.
    reference_top=top+height-37
    answer_top=y+3
    answer_height=reference_top-8-(answer_top+15)
    assert answer_height >= 36, f"Question {q['id']} needs more answer space: {answer_height}"
    para(c,'JAWABAN KAMU / TIM — tulis pilihan + alasan; atau USULKAN / BELUM TAHU',x,answer_top,width,8,TEAL,True)
    field(c,'jawaban_'+q['id'],x,answer_top+15,width,answer_height,True)
    para(c,'REF / LINK / TIMESTAMP — isi "tidak ada" jika tanpa lampiran',x,reference_top,width,7.8,MUTED)
    field(c,'referensi_'+q['id'],x,reference_top+14,width,17)
    layout_log.append({'id':q['id'],'page':page,'content_bottom':y,'answer_top':answer_top})

for s in SECTIONS:
    for part in range(math.ceil(len(s['questions'])/2)):
        start(f"{s['num']}. {s['title']}",
            f"Prioritas: {s['priority']}  |  Pengisi yang disarankan: {s['owner']}",
            f"section-{s['num']}" if part==0 else None)
        for idx,q in enumerate(s['questions'][part*2:part*2+2]):
            question_card(q,111+idx*328)
        if len(s['questions'][part*2:part*2+2])==1:
            box(c,M,454,CW,195,PALE)
            para(c,'CATATAN TAMBAHAN BAGIAN INI',M+13,467,CW-26,9,TEAL,True)
            para(c,'Boleh mencatat perbedaan pendapat, pertanyaan untuk developer, atau batas yang belum tercakup.',M+13,489,CW-26,9.5)
            field(c,'catatan_bagian_'+s['num'],M+13,531,CW-26,100,True)
        para(c,'Status bagian: tulis DRAF / PERLU DISKUSI / FINAL, lalu nama/peran peninjau.',M,775,CW,7.4,MUTED)
        field(c,f"status_{s['num']}_{part}",M+350,789,CW-350,12)

for index in [1,2]:
    start(f'Lembar referensi {index:02d}', 'Duplikasi halaman ini di Word atau kirim file tambahan jika referensinya lebih banyak.',f'reference-{index}')
    para(c,'Gunakan kode yang sama pada kolom REF di setiap pertanyaan terkait.',M,109,CW,10)
    rows=[('Kode referensi (mis. REF-01)', 'kode',150,140),('Untuk pertanyaan (mis. 3.4, 7.2)','untuk',150,0)]
    para(c,'Kode referensi',M,141,160,9,TEAL,True);field(c,f'ref{index}_kode',M,158,150,24)
    para(c,'Nomor pertanyaan terkait',M+169,141,CW-169,9,TEAL,True);field(c,f'ref{index}_untuk',M+169,158,CW-169,24)
    para(c,'Nama file / link / sumber',M,197,CW,9,TEAL,True);field(c,f'ref{index}_file',M,214,CW,29)
    para(c,'Timestamp video / nomor halaman / versi gambar',M,257,CW,9,TEAL,True);field(c,f'ref{index}_time',M,274,CW,24)
    box(c,M,314,CW,172,PALE)
    para(c,'TEMPAT GAMBAR REFERENSI',M+15,333,CW-30,12,TEAL,True)
    para(c,'WORD: tempel screenshot/gambar pada kotak ini.<br/><br/>PDF: gunakan fitur anotasi gambar jika tersedia, atau kirim gambar terpisah dengan nama REF-01-deskripsi.jpg. Kolom PDF ini bukan tombol unggah.<br/><br/>VIDEO / AUDIO / MODEL: cukup tulis nama file/link di atas; file aslinya dikirim terpisah. Jangan mencoba menempel seluruh video ke PDF.',M+15,363,CW-30,10)
    para(c,'ASPEK YANG DISUKAI / PERLU DIIKUTI',M,504,CW,9,TEAL,True);field(c,f'ref{index}_suka',M,521,CW,48,True)
    para(c,'ASPEK YANG TIDAK BOLEH DIIKUTI / BATAS PENERAPAN',M,586,CW,9,TEAL,True);field(c,f'ref{index}_batas',M,603,CW,48,True)
    para(c,'Jenis bahan: referensi saja / aset final / belum tahu. Sumber atau izin jika relevan.',M,669,CW,8.7,MUTED);field(c,f'ref{index}_sumber',M,687,CW,34,True)
    para(c,'Contoh penamaan: REF-01-kamera.jpg + referensi-video.mp4, waktu 00:18-00:20. Satu referensi boleh dipakai untuk beberapa pertanyaan.',M,740,CW,9,MUTED)

start('Ringkasan dan cara mengirim kembali', 'Tidak perlu mengisi jawaban yang sama dua kali. Gunakan nomor pertanyaan untuk meninjau kelengkapan.', 'handoff')
y=110
y=para(c,'Jawaban awal yang paling membantu',M,y,CW,11,bold=True)+10
minimum=[('1.1-1.3','Prioritas referensi, gaya visual, dan suasana.'),('1.5 + 2.1','Setuju/tidak battle contoh; batas perubahan mekanik.'),('3.1-3.4','Acuan bentuk, dua monster, dan ciri wajib.'),('4.1 + 4.3 + 4.4','Aset tersedia, cara mendapatkannya, anggaran jika ada.'),('10.1 + 10.4','HP utama dan prioritas performa.'),('11.1 + 11.3-11.5','Kekurangan utama, tenggat, sasaran versi, PIC keputusan.'),('12.1-12.3','Kriteria berhasil, batas penolakan, dan cara review.')]
for ids,desc in minimum:
    y=para(c,f'<b>{ids}</b> — {desc}',M,y,CW,9.6)+8
y+=9
y=para(c,'Sebelum dikirim',M,y,CW,11,bold=True)+9
for text in ['Setiap pertanyaan prioritas awal sudah dijawab atau diberi USULKAN / BELUM TAHU.',
 'Contoh jawaban tidak dianggap sebagai keputusan tim tanpa dipilih sendiri.',
 'Setiap REF menyebut nomor pertanyaan, file/link, serta aspek yang disukai.',
 'Video yang dirujuk memiliki menit:detik; gambar negatif diberi label JANGAN DIIKUTI.',
 'PDF telah disimpan dan dibuka ulang untuk memeriksa isian; teks panjang disertakan sebagai lanjutan.',
 'Keputusan yang berbeda antaranggota tim masih ditandai DRAF atau PERLU DISKUSI.']:
    y=para(c,'• '+text,M,y,CW,9.3)+7
y+=9
y=para(c,'Paket yang dikirim kembali',M,y,CW,11,bold=True)+8
y=para(c,'1. PDF yang sudah diisi ATAU Word yang sudah diisi.<br/>2. Folder gambar/video/aset dengan kode REF yang sesuai.<br/>3. Catatan jawaban panjang jika tidak muat di formulir.<br/>Nama usulan: Elemonts-Brief-Jawaban-Tim-v1.pdf dan Referensi-Elemonts/.',M,y,CW,9.5)+12
para(c,'CATATAN PENYERAHAN / BAGIAN YANG MASIH PERLU BANTUAN',M,675,CW,9,TEAL,True)
field(c,'catatan_penyerahan',M,695,CW,72,True)
assert page==TOTAL,(page,TOTAL)
c.save()
(ASSETS/'layout-check.json').write_text(json.dumps(layout_log,indent=2))

# Render explanatory visual pages for the Word companion; not screenshots of game UI.
pdf=fitz.open(PDF)
for pn,name in [(3,'guide-reference.png'),(4,'guide-assets.png')]:
    pdf[pn].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(ASSETS/name)

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.6);sec.bottom_margin=Inches(.6)
sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(5)
for nm in ['Title','Heading 1','Heading 2','Heading 3']:
    doc.styles[nm].font.color.rgb=RGBColor.from_string('17333D')
header=sec.header.paragraphs[0];header.text='ELEMONTS  |  FORMULIR KEBUTUHAN VISUAL DAN BATTLE'
header.style='Caption'
footer=sec.footer.paragraphs[0];footer.text='v1.0 • 25 Oktober 2026 • Contoh bukan keputusan final  |  '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def dp(text,boldlabel=None):
    p=doc.add_paragraph()
    if boldlabel:
        p.add_run(boldlabel+': ').bold=True
    p.add_run(text)
    return p

def input_table(label,placeholder,height=None):
    t=doc.add_table(rows=1,cols=1);t.style='Table Grid'
    cell=t.cell(0,0);p=cell.paragraphs[0];p.add_run(label+'\n').bold=True;p.add_run(placeholder)
    if height:
        t.rows[0].height=Inches(height)
    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EFF8F8');cell._tc.get_or_add_tcPr().append(shd)
    return t

doc.add_heading('Elemonts\nFormulir kebutuhan visual dan battle',0)
dp('Versi Word yang dapat diedit — pendamping PDF formulir. 72 pertanyaan dalam 12 bagian.')
dp('Isi hanya kotak JAWABAN KAMU / TIM, kolom REF, dan lembar referensi. Contoh merupakan ilustrasi, bukan keputusan yang sudah diambil.')
input_table('IDENTITAS PENGISI','Nama/peran: [TULIS DI SINI]\nTim: [TULIS DI SINI]\nTanggal/versi jawaban: [TULIS DI SINI]')
doc.add_heading('Acuan yang tetap berlaku',1)
dp('Dokumen konsep dan aturan yang sudah disepakati tetap menjadi acuan. Excel bentuk 37 monster digunakan untuk penampilan, bukan mengganti elemen gameplay. Android serta pilihan bahasa Indonesia/Inggris tetap menjadi arah proyek.')
dp('Pertanyaan melengkapi kebutuhan produksi. Rekomendasi seperti Pyroo dan Croobit sebagai contoh belum menjadi persetujuan otomatis. Jika belum tahu, tulis USULKAN atau BELUM TAHU. Tidak perlu menentukan jumlah poligon, nama tulang, atau detail teknis kamera.')
doc.add_page_break()
doc.add_heading('Arti setiap label',1)
for label,body in labels:
    dp(body,label)
doc.add_heading('Contoh pengisian, bukan keputusan',1)
dp('7.2: B — kamera ringan untuk serangan biasa; close-up hanya skill spesial. REF-01: video Lumi Master, 00:18-00:20. Suka sorotan monster; tidak meminta menyalin karakter.')
dp('Pilihan banyak: 1.1: B > D > C. Belum tahu: USULKAN, dengan batas lancar di HP target. Tidak ada gambar: pakai deskripsi Excel. Tim belum sepakat: tandai DRAF.')
doc.add_heading('Cara memasukkan gambar dan video',1)
dp('Di lembar referensi, ganti [TEMPEL GAMBAR] dengan gambar melalui Insert/Sisipkan > Pictures/Gambar. Tulis kode REF yang sama di pertanyaan terkait. Video/audio/model dikirim sebagai file terpisah; tulis link atau nama file dan timestamp, bukan menempel seluruh video ke dokumen.')
dp('PDF dapat diketik pada pembaca yang mendukung formulir, tetapi tidak memiliki tombol unggah gambar. Word lebih mudah untuk menyusun jawaban panjang dan menempel gambar. Setelah mengisi, simpan salinan dengan nama tim dan nomor versi.')
for nm in ['guide-reference.png','guide-assets.png']:
    doc.add_page_break();doc.add_picture(str(ASSETS/nm),width=Inches(6.55))
for s in SECTIONS:
    for part in range(math.ceil(len(s['questions'])/2)):
        doc.add_page_break()
        doc.add_heading(f"{s['num']}. {s['title']}",1)
        dp(f"Prioritas: {s['priority']} | Pengisi yang disarankan: {s['owner']}")
        for q in s['questions'][part*2:part*2+2]:
            doc.add_heading(q['id']+'  '+q['question'],2)
            for label,key in [('MAKSUD','meaning'),('CARA MENJAWAB / PILIHAN','options'),('CONTOH KASUS + JAWABAN — BUKAN KEPUTUSAN','example'),('LAMPIRAN','evidence')]:
                dp(q[key],label)
            input_table('JAWABAN KAMU / TIM','[TULIS DI SINI: pilihan + alasan; atau USULKAN / BELUM TAHU]',.55)
            input_table('REF / LINK / TIMESTAMP','[Kode REF + file/link + waktu video, atau TIDAK ADA]')
        dp('Status: [DRAF / PERLU DISKUSI / FINAL] | Peninjau: [NAMA/PERAN]')
for idx in [1,2]:
    doc.add_page_break();doc.add_heading(f'Lembar referensi {idx:02d}',1)
    dp('Duplikasi halaman ini jika perlu lebih banyak referensi. Satu REF boleh dipakai untuk beberapa pertanyaan.')
    input_table('IDENTITAS REFERENSI','Kode REF: [TULIS DI SINI]\nNomor pertanyaan: [TULIS DI SINI]\nNama file/link/sumber: [TULIS DI SINI]\nTimestamp/versi: [TULIS DI SINI]')
    input_table('TEMPAT GAMBAR REFERENSI','[TEMPEL GAMBAR DI SINI]\nGanti tulisan ini dengan gambar. Untuk video/audio/model, tulis file/link di atas dan kirim file asli secara terpisah.',2.4)
    input_table('ASPEK YANG DISUKAI / PERLU DIIKUTI','[TULIS DI SINI]',.7)
    input_table('ASPEK YANG TIDAK BOLEH DIIKUTI / BATAS','[TULIS DI SINI]',.7)
    input_table('JENIS DAN SUMBER','[Referensi saja / aset final / belum tahu. Sumber/izin jika relevan.]')
doc.add_page_break();doc.add_heading('Ringkasan dan pengiriman',1)
dp('Tidak perlu mengulang jawaban. Periksa nomor prioritas berikut:')
for ids,desc in minimum:
    dp(ids+' — '+desc)
dp('Kirim Word atau PDF yang sudah diisi, gambar/video dengan kode REF, dan lanjutan jawaban jika ada. Pastikan pilihan yang belum disepakati ditandai DRAF. Buka ulang PDF yang disimpan untuk memeriksa isian; jawaban panjang di kolom bergulir bisa tidak seluruhnya tercetak.')
input_table('CATATAN PENYERAHAN','[Bagian yang belum dijawab, aset yang menyusul, atau bantuan yang diperlukan.]',1)
doc.save(DOCX)
print(json.dumps({'pdf':str(PDF),'pages':page,'questions':len(QUESTIONS),'widgets':sum(len(list(p.widgets() or [])) for p in pdf),'docx':str(DOCX)},indent=2))
