"""Create a plain-text Excel skill brief, without dropdowns or macros.

Export the current skills/roster to JSON first (see accompanying README).
Usage: python scripts/build-skill-template.py /tmp/elemonts-skill-source.json
"""
import json
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/design/Template-Skill-Elemonts.xlsx'
data = json.loads(Path(sys.argv[1]).read_text())
wb = Workbook()
wb.remove(wb.active)
TEAL = '087F80'
NAVY = '17333D'
LINE = 'D5E2E5'
BLUE = 'EAF4FF'
GREEN = 'EDF8F3'
AMBER = 'FFF3DC'
GRAY = 'F2F5F6'
thin = Side(style='thin', color=LINE)
headers = [
    'Kode isian\n(sudah disediakan)',
    'Nama skill\n(isi dahulu)',
    'Monster pengguna\n(isi dahulu)',
    'Elemen\n(boleh USULKAN)',
    'Jenis aksi\n(boleh USULKAN)',
    'Sumber serangan\n(boleh USULKAN)',
    'Gerakan tubuh\n(boleh USULKAN)',
    'Bentuk efek\n(boleh USULKAN)',
    'Reaksi target\n(boleh USULKAN)',
    'Biasa atau khas\n(boleh USULKAN)',
    'Kode referensi\n(opsional)',
    'Catatan tambahan\n(opsional)',
]
widths = [15,24,30,17,22,24,43,43,38,24,25,38]
helptexts = [
    'Kode untuk menghubungkan isian dengan gambar/video. Bukan ID skill di kode game. Boleh dipertahankan.',
    'Nama serangan. Nama sementara juga boleh. Untuk merevisi yang ada, salin nama dari ACUAN SKILL.',
    'Sebutkan satu atau beberapa nama monster. Pisahkan dengan koma; bila banyak, tulis kelompoknya dengan jelas. Lihat DAFTAR MONSTER.',
    'Api, Air, Listrik, Tanah, Netral, Holy, atau Dark. Jika belum tahu, tulis USULKAN. Excel bentuk monster tidak mengganti aturan elemen.',
    'Fisik dekat, tembakan, area, pertahanan, atau pendukung. Fisik = tubuh menyentuh lawan; area = efek menyebar pada arena.',
    'Dari mana serangan berasal? Misalnya mulut, tangan, ekor, tanduk, atau energi di sekitar tubuh. Untuk pertahanan, tulis tubuh/area yang dilindungi.',
    'Apa yang dilakukan tubuh sebelum, saat, dan setelah menyerang? Contoh: menunduk, menerjang, menyundul, lalu kembali ke posisi.',
    'Apa yang terlihat? Contoh: bola api kecil, percikan, gelombang air, atau perisai transparan. Ini tampilan, bukan rumus damage.',
    'Bagaimana lawan bereaksi? Contoh: tersentak, terdorong, atau terpental ringan. Untuk pertahanan, tulis tidak ada target lawan bila sesuai.',
    'Biasa = animasi dasarnya boleh dipakai beberapa skill. Khas = perlu ciri atau urutan aksi khusus. Ini label produksi, bukan otomatis aturan baru.',
    'Isi REF-01, REF-02, dst. Lalu lengkapi sheet REFERENSI. Jika tidak ada gambar/video, tulis TIDAK ADA.',
    'Tulis batas atau permintaan lain. Perubahan aturan (damage, akurasi, PP, status, cara belajar) harus disebut eksplisit; tidak otomatis berubah karena efek visual.',
]

def sheet(name, title, subtitle, cols):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=cols)
    ws.cell(1,1,title).font=Font(name='Calibri',size=18,bold=True,color='FFFFFF')
    ws.cell(1,1).fill=PatternFill('solid',fgColor=TEAL)
    ws.cell(1,1).alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=34
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=cols)
    ws.cell(2,1,subtitle).font=Font(name='Calibri',size=11,color=NAVY)
    ws.cell(2,1).alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[2].height=42
    ws.sheet_properties.tabColor=TEAL
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True
    ws.print_title_rows='1:4'
    ws.oddFooter.center.text='Elemonts | Template skill | &P / &N'
    ws.sheet_view.zoomScale=85
    return ws

def cells(ws,row,values,fill=None,height=65):
    for col,value in enumerate(values,1):
        c=ws.cell(row,col,value)
        c.font=Font(name='Calibri',size=11,color=NAVY)
        c.alignment=Alignment(wrap_text=True,vertical='top')
        c.border=Border(left=thin,right=thin,top=thin,bottom=thin)
        if fill: c.fill=PatternFill('solid',fgColor=fill)
    ws.row_dimensions[row].height=height

def header(ws,labels,row=4):
    cells(ws,row,labels,TEAL,42)
    for c in ws[row]:
        c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
        c.alignment=Alignment(wrap_text=True,vertical='center')

def table(ws,name,last,colcount):
    t=Table(displayName=name,ref=f'A4:{get_column_letter(colcount)}{last}')
    t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,
        showLastColumn=False,showRowStripes=False,showColumnStripes=False)
    ws.add_table(t)
    ws.freeze_panes='D5'

guide=sheet('MULAI DI SINI','ELEMONTS — TEMPLATE DAFTAR SKILL',
    'Tanpa dropdown, tanpa macro. Semua jawaban berupa teks biasa. Mulai dari contoh, lalu isi sheet ISI SKILL. Versi 1 • 25 Oktober 2026.',3)
for i,w in enumerate([29,84,81],1):guide.column_dimensions[get_column_letter(i)].width=w
header(guide,['Topik','Penjelasan sederhana','Contoh / tindakan kamu'])
instructions=[
 ('Tujuan file','Menjelaskan gerakan dan tampilan tiap skill agar animasi sesuai bentuk monster. Tidak perlu memahami istilah teknis 3D.','Satu baris = satu skill untuk pengguna dan gaya gerak yang dijelaskan pada baris itu.'),
 ('Urutan memakai file','1. Baca CONTOH SAJA. 2. Isi ISI SKILL. 3. Jika ada gambar/video, lengkapi REFERENSI.','ACUAN SKILL dan DAFTAR MONSTER membantu melihat data yang sudah ada di proyek.'),
 ('Kolom biru: isi dahulu','Nama skill dan monster pengguna adalah informasi paling penting. Kode isian sudah disediakan.','Mulai dari dua monster percobaan; usulan sebelumnya Pyroo dan Croobit, belum keputusan otomatis.'),
 ('Kolom hijau: boleh minta usulan','Elemen, jenis aksi, sumber, gerakan, efek, reaksi, dan biasa/khas bisa kamu isi jika sudah punya bayangan.','Jika belum tahu, tulis USULKAN atau BELUM TAHU. Tidak harus mengarang penjelasan.'),
 ('Kolom abu: opsional','Referensi dan catatan boleh dikosongkan atau diisi TIDAK ADA bila belum tersedia.','Tidak memiliki gambar bukan penghalang untuk mengisi daftar skill.'),
 ('Arti jenis aksi','Fisik dekat: tubuh mengenai lawan. Tembakan: mengirim efek ke lawan. Area: efek menyebar. Pertahanan: melindungi. Pendukung: membantu kondisi tertentu.','Tulis istilahmu sendiri juga boleh. Nama ini menjelaskan aksi visual dan tidak otomatis mengganti mekanik.'),
 ('Gerakan tubuh vs bentuk efek','Gerakan tubuh = apa yang monster lakukan. Bentuk efek = apa yang muncul di layar.','Tubuh: menarik kepala lalu meludah. Efek: bola api kecil dengan jejak bara.'),
 ('Biasa vs khas','Biasa: dasar animasi boleh dibagi beberapa skill. Khas: punya ciri/urutan khusus untuk monster tertentu.','Skill biasa dapat punya warna berbeda. Skill khas bisa punya pose dan kamera khusus.'),
 ('Jika skill sudah ada','Lihat ACUAN SKILL, lalu salin namanya ke ISI SKILL jika ingin memberi arahan animasi atau revisi.','Catatan: pertahankan fungsi dan PP; ubah gerakan menjadi lebih terasa menghantam.'),
 ('Jika pengguna skill banyak','Boleh menulis beberapa monster dalam satu sel, dipisahkan koma. Jika gerak tiap monster berbeda, buat baris terpisah.','Terjang — Pyroo: menerjang bahu. Terjang — Croobit: menyundul. Dua baris boleh memakai nama skill yang sama.'),
 ('Aturan dan angka','Tidak wajib mengisi damage, akurasi, PP, atau level belajar. Data implementasi saat ini tersedia pada ACUAN SKILL dan bukan otomatis balancing final.','Kalau ingin aturan diubah, tulis jelas pada Catatan tambahan. Efek visual tidak mengubah aturan dengan sendirinya.'),
 ('Gambar / video','Beri kode REF-01, dst. Tulis kode di ISI SKILL, kemudian nama file/link dan waktu adegan di REFERENSI.','REF-01 | serangan.mp4 | 00:12–00:15 | suka ayunan ekor, bukan warna dan karakter.'),
 ('Tempat menempel gambar','Di REFERENSI ada kotak berlabel TEMPEL GAMBAR. Excel: Sisipkan / Insert > Gambar / Pictures. Gambar juga boleh dikirim terpisah.','Nama file: REF-01-ayunan.jpg. Jangan menempel seluruh video ke Excel; kirim video terpisah.'),
 ('Data acuan monster','Nama, elemen gameplay, tahap evolusi, dan kelompok dibaca dari implementasi proyek saat template dibuat.','File bentuk monster tetap khusus tampilan. Jika data gameplay acuan berbeda dari keputusanmu, tulis koreksinya pada Catatan.'),
 ('Selesai mengisi','Simpan sebagai file baru, lalu kirim Excel dan file referensi yang sesuai. 100 baris disediakan; boleh menambah baris.','Nama usulan: Skill-Elemonts-Jawaban-Tim-v1.xlsx. Contoh pada CONTOH SAJA tidak dianggap keputusan tim.'),
]
for r,vals in enumerate(instructions,5):cells(guide,r,vals,'FFFFFF' if r%2 else GREEN,49)
guide.freeze_panes='B5';guide.print_area=f'A1:C{guide.max_row}'

examples=sheet('CONTOH SAJA','CONTOH PENGISIAN — BUKAN KEPUTUSAN GAME',
    'Contoh ini hanya menjelaskan cara menjawab. Skill, gerakan, dan hubungan pengguna di bawah bukan penetapan baru. Isi jawaban asli pada ISI SKILL.',12)
main=sheet('ISI SKILL','ISI DI SINI — DAFTAR SKILL DAN ANIMASINYA',
    'Biru: nama + pengguna diisi dahulu. Hijau: isi bayanganmu atau tulis USULKAN. Abu: referensi/catatan opsional. Satu baris per skill; jika gerak pengguna berbeda, buat baris terpisah.',12)
for ws in [examples,main]:
    for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
    header(ws,headers)
    for i,text in enumerate(helptexts,1):ws.cell(4,i).comment=Comment(text,'Petunjuk Elemonts')
    ws.sheet_view.zoomScale=80
    ws.sheet_properties.tabColor=TEAL if ws==main else 'C1842B'
example_rows=[
 ['CONTOH-01','Semburan Bara (contoh)','Pyroo (contoh saja)','Api','Tembakan','Mulut','Menarik kepala ke belakang, membuka mulut, menyembur, lalu kembali siap.','Semburan api pendek berbentuk kerucut; bara kecil menghilang setelah benturan.','Tubuh tersentak ke belakang, lalu kembali siap.','Biasa','REF-CONTOH-01','Ini contoh format; nama dan animasi belum diputuskan.'],
 ['CONTOH-02','Bertahan (contoh arahan)','Croobit (contoh saja)','Netral','Pertahanan','Energi di sekitar tubuh','Merendahkan tubuh dan bersiap menahan benturan.','Perisai transparan muncul singkat di depan tubuh.','Tidak ada target lawan; pengguna mengambil pose bertahan.','Biasa','TIDAK ADA','Menggambarkan animasi saja. Fungsi dan PP mengikuti aturan yang ada.'],
 ['CONTOH-03','USULKAN NAMA','Pyroo','USULKAN','USULKAN','USULKAN','Ingin gerakan cepat dan lincah; detailnya tolong usulkan.','USULKAN','USULKAN','Khas','TIDAK ADA','Contoh jika hanya punya bayangan umum; belum menetapkan skill khas Pyroo.'],
 ['CONTOH-04','Terjang (contoh revisi)','Croobit (contoh saja)','Netral','Fisik dekat','Kepala / bahu; perlu ditinjau sesuai bentuk','Mendekat, melakukan kontak dengan target, lalu kembali ke posisi.','Debu kecil pada tanah dan kilau benturan singkat.','Mundur satu langkah.','Biasa','TIDAK ADA','Nama skill sudah ada. Tulis permintaan perubahan visual; angka tidak diubah.'],
]
for r,vals in enumerate(example_rows,5):cells(examples,r,vals,AMBER,125)
for r in range(5,105):
    cells(main,r,[f'SK-{r-4:03d}']+[None]*11,height=86)
    for col in range(1,13):main.cell(r,col).fill=PatternFill('solid',fgColor=BLUE if col<=3 else GREEN if col<=10 else GRAY)
table(examples,'ContohSkill',8,12);table(main,'IsianSkill',104,12)
examples.print_area='A1:L8';main.print_area='A1:L14'

ref=sheet('REFERENSI','REFERENSI GAMBAR / VIDEO — OPSIONAL',
    'Isi kode yang sama dengan ISI SKILL. Jika tidak ada referensi, sheet ini boleh dibiarkan. Satu gambar/video boleh dipakai beberapa skill.',7)
for i,w in enumerate([19,25,43,24,52,48,35],1):ref.column_dimensions[get_column_letter(i)].width=w
header(ref,['Kode referensi','Untuk kode isian / skill','Nama file / link','Waktu video / versi','Bagian yang disukai','Bagian yang jangan diikuti','Sumber / keterangan'])
for r in range(5,17):cells(ref,r,[None]*7,height=65)
table(ref,'DaftarReferensi',16,7);ref.freeze_panes='C5'
for idx,start in enumerate([19,30,41],1):
    ref.merge_cells(start_row=start,start_column=1,end_row=start,end_column=7)
    ref.cell(start,1,f'TEMPAT GAMBAR {idx} — tulis kode REF pada baris di atas, lalu tempel gambar di kotak di bawah.').font=Font(size=12,bold=True,color=TEAL)
    ref.row_dimensions[start].height=28
    ref.merge_cells(start_row=start+1,start_column=1,end_row=start+8,end_column=7)
    c=ref.cell(start+1,1,'[TEMPEL GAMBAR DI SINI]\nGunakan Sisipkan / Insert > Gambar / Pictures.\nAtau kirim gambar terpisah dengan nama REF-01-deskripsi.jpg.\nUntuk video/audio/model: tulis file/link dan kirim file aslinya terpisah.')
    c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    c.font=Font(size=13,color=TEAL);c.fill=PatternFill('solid',fgColor=GREEN)
    for rr in range(start+1,start+9):ref.row_dimensions[rr].height=25
ref.print_area='A1:G49'

current=sheet('ACUAN SKILL','24 SKILL YANG SUDAH ADA DI PROYEK',
    'ACUAN, bukan kolom wajib diisi dan bukan balancing final. Data diambil dari src/game/data.ts. Salin nama ke ISI SKILL bila ingin memberi arahan. Power bukan damage akhir.',10)
for i,w in enumerate([18,26,17,60,18,19,16,20,35,50],1):current.column_dimensions[get_column_letter(i)].width=w
header(current,['ID di proyek','Nama skill','Elemen','Pengguna menurut implementasi','Power dasar','Akurasi (%)','PP awal','Level belajar','Efek yang ada','Keterangan'])
for r,s in enumerate(data['skills'],5):
    effects={'burn':'Burn (terbakar)','slow':'Slow (melambat)','guard':'Guard (bertahan)'}
    note='Aksi cadangan otomatis; bukan skill yang dipelajari/dipasang normal.' if s['id']=='struggle' else 'Data implementasi saat template dibuat; fungsi final mengikuti keputusan desain.'
    vals=[s['id'],s['name'],s['element'],data['users'][s['id']],s['power'],s['accuracy'],s['pp'],None if s['id']=='struggle' else s['learn'],effects.get(s.get('effect'),'Tidak ada efek tambahan'),note]
    cells(current,r,vals,GRAY,112 if len(vals[3])>200 else 76)
table(current,'SkillImplementasi',28,10);current.print_area='A1:J28'

roster=sheet('DAFTAR MONSTER','37 MONSTER — ACUAN NAMA DAN ELEMEN GAMEPLAY',
    'Tidak perlu menulis ulang daftar ini. Pakai namanya pada ISI SKILL. Acuan berasal dari implementasi, bukan pengganti keputusan konsep atau Excel bentuk monster.',6)
for i,w in enumerate([15,27,22,20,25,30],1):roster.column_dimensions[get_column_letter(i)].width=w
header(roster,['ID proyek','Nama monster','Elemen gameplay','Tahap evolusi','Kategori','Keluarga evolusi'])
for r,s in enumerate(data['roster'],5):cells(roster,r,[s['id'],s['name'],s['element'],s['stage'],s['category'],s['family']],GRAY,27)
table(roster,'AcuanMonster',41,6);roster.print_area='A1:F41'

for ws in wb:
    ws.sheet_properties.pageSetUpPr.autoPageBreaks=False
    ws.sheet_view.tabSelected=False
wb.active=0
wb.save(OUT)

# Reopen the deliverable to verify that it is a blank, plain-text template.
check=load_workbook(OUT)
assert len(data['skills'])==24 and len(data['roster'])==37
assert check.sheetnames==['MULAI DI SINI','CONTOH SAJA','ISI SKILL','REFERENSI','ACUAN SKILL','DAFTAR MONSTER']
assert all(check['ISI SKILL'].cell(r,c).value is None for r in range(5,105) for c in range(2,13))
assert all(len(ws.data_validations.dataValidation)==0 for ws in check)
assert not any(c.data_type=='f' for ws in check for row in ws for c in row)
assert check['ACUAN SKILL'].max_row==28 and check['DAFTAR MONSTER'].max_row==41
print(json.dumps({'path':str(OUT),'sheets':check.sheetnames,'blank_input_rows':100,'existing_skills':24,'monsters':37,'dropdowns':0},ensure_ascii=False,indent=2))
